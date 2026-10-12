"""เขียนผลรันชุดกันถอยหลังลง docs/tester.xlsx (แผ่นเดียว "All Test Cases & Results").

- ข้อเดิมที่อยู่ในไฟล์: อัปเดตคำตอบ หน้าอ้างอิง เส้นทาง และเวลาเป็นผลรอบล่าสุด
  ผลตรวจโดยคนคงไว้ เว้นแต่ข้อที่ตรวจซ้ำแล้ว (RECHECKED) ซึ่งมีหมายเหตุกำกับ
- ข้อใหม่ (คลังคำถาม / ชุดที่ไม่เคยเห็น): ต่อท้าย ผลตรวจมาจากการตรวจอัตโนมัติเทียบเฉลย

Usage:
    python -m katrag.eval.export_tester artifacts/regression_r10.json
"""

from __future__ import annotations

import json
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

ROOT = Path(__file__).resolve().parents[2]
TESTER = ROOT / "docs" / "tester.xlsx"

#: ข้อเดิมที่ตรวจคำตอบซ้ำหลังปรับระบบ (อ่านคำตอบทั้งหมดเทียบเฉลยแล้ว)
RECHECKED = {
    "T10": ("ถูกครบ", "ตรวจซ้ำหลังปรับระบบ: ตอบครบ 3 วุฒิ (วศ.บ. สจล. 2523, M.Eng AIT 2525, Ph.D. "
                      "The University of Queensland 2531) จากตารางประวัติอาจารย์ AIT 2566 หน้า 42 ไม่มีเลขประจำตัวประชาชน"),
    "X5": ("ถูกครบ", "ตรวจซ้ำหลังปรับระบบ: 2563 แผน ก (ก1, ก2) และแผน ข; 2568 แผน 1 แบบวิชาการ (1.1, 1.2) "
                     "และแผน 2 แบบวิชาชีพ ทุกแผน 36 หน่วยกิต อ่านจากหัวข้อ 3.1.1 ของทั้งสองเล่ม"),
    # สองข้อนี้ในไฟล์เดิมบันทึกเป็น "ถูกครบ" พร้อมหมายเหตุของข้ออื่น (แถวเลื่อนมาจากฉบับเก่า) — ตัดสินใหม่ตามคำตอบล่าสุด
    "H2": ("ถูกบางส่วน", "ตรวจซ้ำ: ตอบ 'ได้' ถูก แต่ตัวอย่างวิชาปี 4 ไม่ตรงเฉลย (ระบบฐานข้อมูลแบบกระจาย) "
                         "ข้อนี้ผ่าน LLM คำตอบเปลี่ยนได้ทุกรอบ"),
    "C1": ("ถูกบางส่วน", "ตรวจซ้ำ: หน่วยกิตรวม 130 → 129 ถูก ส่วนโครงสร้างตอบว่าเปลี่ยนจากแขนงวิชาเป็นโมดูลอาชีพ 3 โมดูล "
                         "ซึ่งตรงกับเล่ม IT 2565 แต่ไม่ตรงกับเฉลยที่ระบุ 3 แขนง"),
}

LEVEL = {"ง่าย": "ระดับ 1: ง่าย", "ปานกลาง": "ระดับ 2: ปานกลาง", "ยาก": "ระดับ 3: ยาก", "ท้าทาย": "ระดับ 4: ท้าทาย"}

#: ชนิดคำถามในคลัง → (ระดับ, ทักษะที่วัด)
BANK_KIND = {
    "credit": ("ง่าย", "หน่วยกิตของวิชาจากชื่อไทย"),
    "credit_para": ("ง่าย", "หน่วยกิตของวิชา (สำนวนอื่น)"),
    "en_credit": ("ง่าย", "หน่วยกิตของวิชาจากชื่ออังกฤษ"),
    "en_name": ("ง่าย", "ชื่อภาษาอังกฤษของวิชา"),
    "by_code": ("ง่าย", "ชื่อวิชาจากรหัสวิชา"),
    "gened": ("ง่าย", "หน่วยกิตวิชาศึกษาทั่วไป"),
    "yearsem": ("ง่าย", "ปี/ภาคที่เรียนของวิชา"),
    "en_yearsem": ("ปานกลาง", "ปี/ภาคที่เรียน ถามด้วยชื่ออังกฤษ"),
    "code_ys": ("ปานกลาง", "ปี/ภาคที่เรียน ถามด้วยรหัสวิชา (ยึดหลักสูตรที่เลือก)"),
    "typo_ys": ("ปานกลาง", "ปี/ภาคที่เรียน เมื่อพิมพ์ชื่อวิชาผิด"),
    "prereq": ("ปานกลาง", "วิชาบังคับก่อน"),
    "semester": ("ปานกลาง", "รายวิชาบังคับทั้งภาคเรียน"),
    "dependents": ("ยาก", "วิชาบังคับก่อนทิศกลับ (ไม่ผ่านแล้วลงอะไรไม่ได้)"),
    "coop": ("ยาก", "แผนสหกิจศึกษา: ปี/ภาคและหน่วยกิต"),
    "rule": ("ยาก", "เกณฑ์จบ/เกรด/เกียรตินิยม/การลงทะเบียน"),
    "reg2": ("ยาก", "ข้อบังคับปริญญาตรี (สำนวนอื่น)"),
}

PATH = {
    "year_sem": "Structured SQL (ปี/ภาค)", "prerequisite": "Structured SQL (prerequisite)",
    "dependents": "Structured SQL (prerequisite ทิศกลับ)", "course_lookup": "Structured SQL (ค้นวิชา: ตารางแผนและหัวข้อรายวิชา)",
    "course_code": "Structured SQL (รหัสวิชา)", "rule": "Structured SQL (ตาราง rule)",
    "category_credits": "Structured SQL (หน่วยกิตรายหมวด)", "credit_compare": "Structured SQL (เทียบหน่วยกิตข้ามเวอร์ชัน)",
    "plan_variant": "Structured SQL (ตารางแผนปกติ/แผนสหกิจ)", "semester_load": "Structured SQL (หน่วยกิตต่อภาค)",
    "track": "Structured SQL (วิชาในแขนง/กลุ่มวิชา)", "person": "Structured SQL (รายชื่ออาจารย์)",
    "person_education": "ประวัติอาจารย์ (ตารางหัวข้อ 3.2)", "regulation": "ข้อบังคับปริญญาตรี (ภาคผนวกเล่ม AIT 2566)",
    "grad_plans": "แผนการศึกษาบัณฑิตศึกษา (หัวข้อ 3.1.1/3.1.2)", "plan_summary": "Structured SQL (สรุปแผนการเรียน)",
    "cross_version": "Structured SQL (เทียบรายวิชาข้ามเวอร์ชัน)", "program_name": "ชื่อหลักสูตร (หน้าข้อมูลทั่วไป)",
    "topic_semantic": "Semantic topic search (bge-m3) → จัดรูปคำตอบเอง",
    "can_take": "Multi-hop SQL (ปีตามแผน + วิชาบังคับก่อน → ลงได้/ไม่ได้)",
    "course_compare": "Multi-hop SQL (เทียบปี/ภาคของสองวิชา)",
    "cross_program_course": "Multi-hop SQL (วิชาเดียวกันข้ามหลักสูตร)",
    "prereq_chain": "Multi-hop SQL (วิชาบังคับก่อนไล่หลายชั้น)",
    "grad_check": "Multi-hop SQL (หน่วยกิตรวม + เกรดขั้นต่ำ ข้อ 34)",
    "program_compare": "Multi-hop SQL (หน่วยกิตรวมข้ามหลักสูตร)",
    "semester_total": "Multi-hop SQL (หน่วยกิตรวมของภาคตามแผน)",
}


def path_label(r: dict) -> str:
    base = PATH.get(r.get("intent") or "", "")
    if not base:
        return "Hybrid retrieval (คำสำคัญ + bge-m3, RRF) → Typhoon" if r.get("llm_calls") else "Hybrid retrieval"
    return base + (" → Typhoon เรียบเรียง" if r.get("llm_calls") else " — ไม่ผ่าน LLM")


def main(results_path: str) -> None:
    res = {r["id"]: r for r in json.loads(Path(results_path).read_text(encoding="utf-8"))}
    meta = {}
    for name in ("eval_bank.json", "eval_holdout.json", "eval_holdout2.json", "eval_holdout3.json", "eval_holdout4.json",
                 "eval_holdout5.json", "eval_holdout6.json", "eval_holdout7.json"):
        p = ROOT / "docs" / name
        if p.exists():
            for it in json.loads(p.read_text(encoding="utf-8"))["items"]:
                meta[it["id"]] = it

    wb = openpyxl.load_workbook(TESTER)
    ws = wb["All Test Cases & Results"]
    style_row = [ws.cell(2, c) for c in range(1, 16)]
    existing = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}

    def put(row: int, r: dict) -> None:
        pages = ", ".join(str(p) for p in r.get("pages") or [])
        ws.cell(row, 7).value = r["answer"]
        ws.cell(row, 8).value = f"หน้า {pages}" if pages else "-"
        ws.cell(row, 9).value = path_label(r)
        ws.cell(row, 10).value = len(r.get("pages") or [])
        ws.cell(row, 11).value = "PASS" if r["status"] in ("SAME", "PASS") else r["status"]
        ws.cell(row, 14).value = f"{r['latency']:.2f}s"
        ws.cell(row, 15).value = "PASS" if r["latency"] < 5 else "FAIL"

    for rid, row in existing.items():
        if rid in res:
            put(row, res[rid])
        if rid in RECHECKED:
            ws.cell(row, 12).value, ws.cell(row, 13).value = RECHECKED[rid]

    row = ws.max_row + 1
    order = [r for r in res.values() if r["set"] in ("bank", "holdout", "holdout2", "holdout3", "holdout4", "holdout5", "holdout6", "holdout7") and r["id"] not in existing]
    for r in order:
        m = meta.get(r["id"], {})
        if r["set"] == "bank":
            kind = r["id"].split("-")[2]
            level, skill = BANK_KIND.get(kind, ("ปานกลาง", kind))
            group = "คลังคำถามจากเฉลยอาจารย์"
        else:
            level = m.get("level", "ปานกลาง")
            skill = "ชุดที่ระบบไม่เคยเห็น" + {"holdout2": " รอบ 2", "holdout3": " รอบ 3", "holdout4": " รอบ 4",
                                          "holdout5": " รอบ 5 (ยาก/ท้าทาย)", "holdout6": " รอบ 6 (ยาก/ท้าทาย)", "holdout7": " รอบ 7 (ยาก/ท้าทาย แบบใหม่)"}.get(r["set"], "")
            group = skill
        passed = r["status"] in ("SAME", "PASS")
        values = [
            r["id"], LEVEL.get(level, level), r["program"], r["question"], f"{skill} — {group}",
            m.get("expected", ""), None, None, None, None, None,
            "ถูกครบ" if passed else "ไม่ผ่าน",
            ("ตรวจอัตโนมัติ: คำตอบมีค่าตามเฉลยครบ" if passed else "ตรวจอัตโนมัติ: " + "; ".join(r.get("problems") or []))
            + f" (เฉลยจาก {m.get('source', '')})",
            None, None,
        ]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row, c)
            cell.value = v
            src = style_row[c - 1]
            cell.font, cell.border, cell.fill = copy(src.font), copy(src.border), copy(src.fill)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        put(row, r)
        row += 1
    wb.save(TESTER)
    print(f"อัปเดต {len(existing)} ข้อเดิม, เพิ่ม {len(order)} ข้อใหม่ -> {TESTER}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(ROOT / "artifacts" / "regression_last.json"))
